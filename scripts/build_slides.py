"""
Builds the project presentation as a .pptx file.
Upload to Google Drive → it auto-converts to Google Slides.
"""
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

ROOT = Path(__file__).parent.parent
EVAL = json.load(open(ROOT / "scripts" / "eval_results.json"))

# ── Brand colours ──────────────────────────────────────────────────────────────
NAVY   = RGBColor(0x00, 0x30, 0x87)   # UMass Boston blue
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT  = RGBColor(0xE8, 0xF0, 0xFF)
GOLD   = RGBColor(0xF5, 0xC5, 0x18)
GREEN  = RGBColor(0x16, 0xA3, 0x4A)
RED    = RGBColor(0xDC, 0x26, 0x26)
GRAY   = RGBColor(0x6B, 0x72, 0x80)

W = Inches(13.33)
H = Inches(7.5)


def new_prs():
    prs = Presentation()
    prs.slide_width  = W
    prs.slide_height = H
    return prs


def blank_slide(prs):
    blank_layout = prs.slide_layouts[6]
    return prs.slides.add_slide(blank_layout)


def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def box(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def text_box(slide, text, left, top, width, height, size, bold=False,
             color=WHITE, align=PP_ALIGN.LEFT, wrap=True):
    txb = slide.shapes.add_textbox(left, top, width, height)
    txb.word_wrap = wrap
    tf = txb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return txb


def bullet_box(slide, items, left, top, width, height, size=16,
               color=RGBColor(0x1F, 0x29, 0x37), title=None, title_size=18):
    txb = slide.shapes.add_textbox(left, top, width, height)
    txb.word_wrap = True
    tf = txb.text_frame
    tf.word_wrap = True
    first = True
    if title:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = title
        r.font.size = Pt(title_size)
        r.font.bold = True
        r.font.color.rgb = NAVY
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.level = 0
        r = p.add_run()
        r.text = f"• {item}"
        r.font.size = Pt(size)
        r.font.color.rgb = color


# ──────────────────────────────────────────────────────────────────────────────
# SLIDE 1 — Title
# ──────────────────────────────────────────────────────────────────────────────
def slide_title(prs):
    s = blank_slide(prs)
    bg(s, NAVY)
    box(s, 0, Inches(5.8), W, Inches(1.7), RGBColor(0x00, 0x20, 0x60))
    text_box(s, "EZRide", Inches(0.7), Inches(1.0), Inches(12), Inches(1.4),
             54, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    text_box(s, "UMass Boston Transportation Assistant",
             Inches(0.7), Inches(2.3), Inches(12), Inches(0.8),
             26, color=GOLD, align=PP_ALIGN.CENTER)
    text_box(s, "RAG-Powered Chatbot with Interactive Campus Map",
             Inches(0.7), Inches(3.1), Inches(12), Inches(0.6),
             18, color=LIGHT, align=PP_ALIGN.CENTER)
    text_box(s, "CS Project Presentation  ·  UMass Boston  ·  2026",
             Inches(0.7), Inches(6.0), Inches(12), Inches(0.5),
             14, color=LIGHT, align=PP_ALIGN.CENTER)


# ──────────────────────────────────────────────────────────────────────────────
# SLIDE 2 — Problem & Use Case
# ──────────────────────────────────────────────────────────────────────────────
def slide_problem(prs):
    s = blank_slide(prs)
    bg(s, WHITE)
    box(s, 0, 0, W, Inches(1.1), NAVY)
    text_box(s, "1. Problem & Use Case", Inches(0.5), Inches(0.15),
             Inches(12), Inches(0.8), 28, bold=True, color=WHITE)

    # Problem card
    box(s, Inches(0.5), Inches(1.3), Inches(5.8), Inches(5.6), LIGHT)
    text_box(s, "The Problem", Inches(0.7), Inches(1.45), Inches(5.4), Inches(0.5),
             18, bold=True, color=NAVY)
    bullet_box(s, [
        "UMass Boston's transportation info is scattered across 11+ web pages",
        "Students waste time searching for permit rules, shuttle times, and MBTA passes",
        "No single interface combines navigation + Q&A",
        "Existing pages have no conversational interface — just static text",
    ], Inches(0.7), Inches(1.95), Inches(5.4), Inches(4.5), size=15,
        color=RGBColor(0x1F, 0x29, 0x37))

    # Users card
    box(s, Inches(6.9), Inches(1.3), Inches(5.9), Inches(5.6), RGBColor(0xF0, 0xF9, 0xF0))
    text_box(s, "Target Users", Inches(7.1), Inches(1.45), Inches(5.5), Inches(0.5),
             18, bold=True, color=NAVY)
    bullet_box(s, [
        "Students — finding parking, transit passes, bike facilities",
        "Faculty & Staff — permit purchases, pre-tax benefits, carpool subsidy",
        "Visitors — daily parking rates and campus navigation",
        "New students — unfamiliar with campus layout and services",
    ], Inches(7.1), Inches(1.95), Inches(5.5), Inches(4.5), size=15,
        color=RGBColor(0x1F, 0x29, 0x37))


# ──────────────────────────────────────────────────────────────────────────────
# SLIDE 3 — System Overview
# ──────────────────────────────────────────────────────────────────────────────
def slide_system(prs):
    s = blank_slide(prs)
    bg(s, WHITE)
    box(s, 0, 0, W, Inches(1.1), NAVY)
    text_box(s, "2. System Overview", Inches(0.5), Inches(0.15),
             Inches(12), Inches(0.8), 28, bold=True, color=WHITE)

    # Pipeline boxes
    stages = [
        ("Web Scraper", "BeautifulSoup\n11 UMass pages", RGBColor(0xDB, 0xEA, 0xFE)),
        ("Chunker", "400-word chunks\n80-word overlap\n44 total chunks", RGBColor(0xDC, 0xFA, 0xE8)),
        ("Embedder", "all-MiniLM-L6-v2\n384-dim vectors\nFAISS index", RGBColor(0xFE, 0xF3, 0xC7)),
        ("Hybrid Search", "BM25 + FAISS\nRRF fusion\n75% semantic", RGBColor(0xFE, 0xE2, 0xE2)),
        ("Gemini LLM", "gemini-3.1-flash-lite\nGevent streaming\nThinking disabled", RGBColor(0xF3, 0xE8, 0xFF)),
    ]
    arrow_w = Inches(0.3)
    box_w   = Inches(2.1)
    gap     = Inches(0.18)
    start_x = Inches(0.35)
    top     = Inches(1.3)
    box_h   = Inches(1.8)

    for i, (title, detail, color) in enumerate(stages):
        x = start_x + i * (box_w + arrow_w + gap)
        box(s, x, top, box_w, box_h, color)
        text_box(s, title, x + Inches(0.1), top + Inches(0.1),
                 box_w - Inches(0.2), Inches(0.4), 13, bold=True,
                 color=NAVY, align=PP_ALIGN.CENTER)
        text_box(s, detail, x + Inches(0.1), top + Inches(0.55),
                 box_w - Inches(0.2), Inches(1.1), 11,
                 color=RGBColor(0x1F, 0x29, 0x37), align=PP_ALIGN.CENTER)
        if i < len(stages) - 1:
            ax = x + box_w + gap * 0.2
            text_box(s, "→", ax, top + Inches(0.65), arrow_w, Inches(0.5),
                     22, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

    # Tools / APIs grid
    text_box(s, "Tools, APIs & Models", Inches(0.5), Inches(3.35),
             Inches(12), Inches(0.4), 16, bold=True, color=NAVY)

    tools = [
        ("Frontend", "Leaflet.js map, Stadia Maps tiles, OSRM routing"),
        ("Walking Routes", "Stadia Maps Valhalla (via Vercel proxy)"),
        ("Transit Data", "MBTA API, UMass TransLoc arrivals"),
        ("Search", "rank-bm25 + FAISS (faiss-cpu)"),
        ("LLM", "Google Gemini 3.1 Flash Lite"),
        ("Hosting", "HF Spaces (backend) + Vercel (frontend)"),
    ]
    col_w = Inches(4.0)
    for i, (label, val) in enumerate(tools):
        col = i % 3
        row = i // 3
        x = Inches(0.5) + col * (col_w + Inches(0.4))
        y = Inches(3.85) + row * Inches(0.95)
        box(s, x, y, col_w, Inches(0.8), LIGHT)
        text_box(s, label, x + Inches(0.1), y + Inches(0.05),
                 col_w - Inches(0.2), Inches(0.3), 11, bold=True, color=NAVY)
        text_box(s, val, x + Inches(0.1), y + Inches(0.38),
                 col_w - Inches(0.2), Inches(0.35), 10,
                 color=RGBColor(0x1F, 0x29, 0x37))


# ──────────────────────────────────────────────────────────────────────────────
# SLIDE 4 — What Works & What Doesn't
# ──────────────────────────────────────────────────────────────────────────────
def slide_cases(prs):
    s = blank_slide(prs)
    bg(s, WHITE)
    box(s, 0, 0, W, Inches(1.1), NAVY)
    text_box(s, "3. What Works & What Doesn't", Inches(0.5), Inches(0.15),
             Inches(12), Inches(0.8), 28, bold=True, color=WHITE)

    # Success
    box(s, Inches(0.5), Inches(1.25), Inches(5.8), Inches(5.7),
        RGBColor(0xF0, 0xFD, 0xF4))
    box(s, Inches(0.5), Inches(1.25), Inches(5.8), Inches(0.5), GREEN)
    text_box(s, "✓  Successful Case", Inches(0.6), Inches(1.3),
             Inches(5.6), Inches(0.4), 14, bold=True, color=WHITE)

    text_box(s, 'Q: "How much does an annual Blue Bikes membership cost?"',
             Inches(0.65), Inches(1.9), Inches(5.5), Inches(0.6), 12,
             bold=True, color=NAVY)
    text_box(s, 'A (RAG): "The annual Bluebikes membership for UMass Boston affiliates '
               'costs $101.50 using promo code BikeUMB."',
             Inches(0.65), Inches(2.55), Inches(5.5), Inches(0.8), 12,
             color=RGBColor(0x1F, 0x29, 0x37))

    text_box(s, "Why it worked:", Inches(0.65), Inches(3.45),
             Inches(5.5), Inches(0.3), 12, bold=True, color=NAVY)
    bullet_box(s, [
        "Exact price $101.50 and promo code BikeUMB present in biking chunk",
        "Semantic search correctly prioritizes the biking page",
        "Short, factual question — single chunk is sufficient",
        "Scored 5/5 on all 5 dimensions",
    ], Inches(0.65), Inches(3.78), Inches(5.5), Inches(2.8), size=12,
        color=RGBColor(0x1F, 0x29, 0x37))

    # Challenge
    box(s, Inches(6.9), Inches(1.25), Inches(5.9), Inches(5.7),
        RGBColor(0xFF, 0xF7, 0xED))
    box(s, Inches(6.9), Inches(1.25), Inches(5.9), Inches(0.5), RGBColor(0xD9, 0x77, 0x06))
    text_box(s, "~  Challenge: Vague Queries", Inches(7.0), Inches(1.3),
             Inches(5.7), Inches(0.4), 14, bold=True, color=WHITE)

    text_box(s, 'Q: "Tell me about permits"',
             Inches(7.05), Inches(1.9), Inches(5.65), Inches(0.4), 12,
             bold=True, color=NAVY)
    text_box(s, 'A (before fix): "I don\'t have enough specific information to answer that."',
             Inches(7.05), Inches(2.35), Inches(5.65), Inches(0.6), 12,
             color=RED)
    text_box(s, 'A (after fix): Full summary of permit types, portal steps, rules, and cancellation.',
             Inches(7.05), Inches(2.95), Inches(5.65), Inches(0.6), 12,
             color=GREEN)

    text_box(s, "Root cause & fix:", Inches(7.05), Inches(3.7),
             Inches(5.65), Inches(0.3), 12, bold=True, color=NAVY)
    bullet_box(s, [
        "Retrieval was correct — 5 permit chunks always returned",
        "LLM treated vague query like a missing specific fact → refused",
        "Fix: system prompt now distinguishes broad vs specific questions",
        "Broad → synthesize all retrieved context; Specific → refuse if missing",
    ], Inches(7.05), Inches(4.05), Inches(5.65), Inches(2.5), size=12,
        color=RGBColor(0x1F, 0x29, 0x37))


# ──────────────────────────────────────────────────────────────────────────────
# SLIDE 5 — Evaluation
# ──────────────────────────────────────────────────────────────────────────────
def slide_eval(prs):
    s = blank_slide(prs)
    bg(s, WHITE)
    box(s, 0, 0, W, Inches(1.1), NAVY)
    text_box(s, "4. Evaluation", Inches(0.5), Inches(0.15),
             Inches(12), Inches(0.8), 28, bold=True, color=WHITE)

    # Method
    text_box(s, "Method", Inches(0.5), Inches(1.2), Inches(12), Inches(0.35),
             14, bold=True, color=NAVY)
    text_box(s,
             "10 factual questions grounded in the corpus were run through (1) the full RAG pipeline "
             "and (2) the Gemini LLM with no retrieval context. Both answers were scored 1–5 on five "
             "dimensions by Gemini acting as an automated judge.",
             Inches(0.5), Inches(1.55), Inches(12.3), Inches(0.65), 13,
             color=RGBColor(0x1F, 0x29, 0x37))

    # Score table
    dims = ["Accuracy", "Completeness", "Relevance", "Conciseness", "Groundedness"]
    rag_scores  = [5.00, 5.00, 4.80, 4.20, 5.00]
    dir_scores  = [2.30, 1.30, 4.10, 4.20, 4.20]

    table_left = Inches(0.5)
    table_top  = Inches(2.35)
    col_widths = [Inches(2.8), Inches(2.0), Inches(2.0)]
    row_h      = Inches(0.48)
    headers    = ["Dimension", "RAG Pipeline", "Direct LLM"]

    # header row
    for ci, (hdr, cw) in enumerate(zip(headers, col_widths)):
        cx = table_left + sum(col_widths[:ci])
        box(s, cx, table_top, cw, row_h, NAVY)
        text_box(s, hdr, cx + Inches(0.08), table_top + Inches(0.08),
                 cw - Inches(0.1), row_h - Inches(0.1), 13, bold=True,
                 color=WHITE, align=PP_ALIGN.CENTER)

    for ri, (dim, rs, ds) in enumerate(zip(dims, rag_scores, dir_scores)):
        y = table_top + (ri + 1) * row_h
        row_bg = LIGHT if ri % 2 == 0 else WHITE
        cx0 = table_left
        box(s, cx0, y, col_widths[0], row_h, row_bg)
        text_box(s, dim, cx0 + Inches(0.08), y + Inches(0.08),
                 col_widths[0] - Inches(0.1), row_h - Inches(0.1), 13,
                 color=RGBColor(0x1F, 0x29, 0x37))
        # RAG col
        cx1 = table_left + col_widths[0]
        rag_c = GREEN if rs >= ds else RED
        box(s, cx1, y, col_widths[1], row_h, row_bg)
        text_box(s, f"{rs:.2f} / 5", cx1, y + Inches(0.08),
                 col_widths[1], row_h - Inches(0.1), 14, bold=True,
                 color=rag_c, align=PP_ALIGN.CENTER)
        # Direct col
        cx2 = table_left + col_widths[0] + col_widths[1]
        dir_c = GREEN if ds >= rs else RED
        box(s, cx2, y, col_widths[2], row_h, row_bg)
        text_box(s, f"{ds:.2f} / 5", cx2, y + Inches(0.08),
                 col_widths[2], row_h - Inches(0.1), 14, bold=True,
                 color=dir_c, align=PP_ALIGN.CENTER)

    # Insights
    text_box(s, "Key Findings", Inches(7.3), Inches(2.25),
             Inches(5.5), Inches(0.4), 15, bold=True, color=NAVY)
    bullet_box(s, [
        "RAG scores 5.00/5 on accuracy vs 2.30 direct — retrieval prevents hallucination",
        "RAG scores 5.00/5 on completeness vs 1.30 direct — context drives full answers",
        "Both score similarly on conciseness — RAG doesn't pad answers unnecessarily",
        "7/10 questions had all key terms in retrieved top-5 chunks",
        "3/10 missed retrieval hints but LLM recovered from nearby context",
        "Direct LLM hallucinated wrong prices, codes, and locations on 7/10 questions",
    ], Inches(7.3), Inches(2.7), Inches(5.7), Inches(4.3), size=13,
        color=RGBColor(0x1F, 0x29, 0x37))


# ──────────────────────────────────────────────────────────────────────────────
# BUILD
# ──────────────────────────────────────────────────────────────────────────────
def main():
    prs = new_prs()
    slide_title(prs)
    slide_problem(prs)
    slide_system(prs)
    slide_cases(prs)
    slide_eval(prs)

    out = ROOT / "EZRide_Presentation.pptx"
    prs.save(str(out))
    print(f"Saved → {out}")


if __name__ == "__main__":
    main()
